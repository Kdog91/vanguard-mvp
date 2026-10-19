"""
One-click Word report: what was loaded, what the cleaner changed, and a plain-language
description of every column, with charts. Built with python-docx and matplotlib so it
works offline and opens in Word, Google Docs or LibreOffice.
"""
import io
from datetime import date

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

import descriptive as d

NAVY = RGBColor(0x14, 0x26, 0x3D)
BAR = "#1f5fb0"
MAX_NUMERIC = 6       # columns described in detail, to keep the report readable
MAX_CATEGORICAL = 6


def _money_like(col: str) -> bool:
    c = col.lower()
    return any(k in c for k in ("amount", "cost", "price", "spend", "obligat", "outlay", "value", "dollar", "revenue"))


def _fmt(v, col: str = "") -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "–"
    if isinstance(v, (int, np.integer)) or (isinstance(v, float) and float(v).is_integer() and abs(v) >= 1000):
        s = f"{v:,.0f}"
    elif isinstance(v, float):
        s = f"{v:,.2f}"
    else:
        s = str(v)
    return ("$" + s) if _money_like(col) and isinstance(v, (int, float, np.number)) else s


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _numeric_chart(s: pd.Series, col: str):
    s = s.dropna().astype(float)
    fig, ax = plt.subplots(figsize=(6.2, 2.4))
    log = abs(float(s.skew())) > 2 and s.min() >= 0 and (s > 0).sum() > 2
    if log:
        pos = s[s > 0]
        bins = np.logspace(np.log10(pos.min()), np.log10(pos.max()), 25)
        ax.hist(pos, bins=bins, color=BAR)
        ax.set_xscale("log")
        ax.set_xlabel(f"{col} (log scale)")
    else:
        ax.hist(s, bins=25, color=BAR)
        ax.set_xlabel(col)
    ax.set_ylabel("Rows")
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    return _png(fig), log


def _categorical_chart(counts: pd.Series, col: str):
    top = counts.head(10)[::-1]
    fig, ax = plt.subplots(figsize=(6.2, 0.3 * len(top) + 0.6))
    ax.barh([str(x)[:40] for x in top.index], top.values, color=BAR)
    ax.set_xlabel("Rows")
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    return _png(fig)


def build_report(dataset_label: str, source_line: str, raw_df: pd.DataFrame, clean_df: pd.DataFrame,
                 profile: dict, cleaning_lines: list, metrics: dict, notes: list | None = None) -> bytes:
    """Return a .docx report as bytes.
    metrics: {"rows_after", "cols_kept", "cols_raw", "rows_removed", "filled", "flagged"}
    notes: extra plain-language warnings (for example, date problems found in the source)."""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)

    h = doc.add_heading("Data Analysis Report", level=0)
    h.runs[0].font.color.rgb = NAVY
    p = doc.add_paragraph()
    p.add_run(dataset_label.replace("_", " ")).bold = True
    doc.add_paragraph(f"{source_line}\nPrepared {date.today():%B %d, %Y} with GovData Analytics, "
                      "Vanguard Data Analytics LLC.")

    # ---- summary ----
    doc.add_heading("1. Summary", level=1)
    num_cols = [c for c, t in profile.items() if t == "numeric"]
    cat_cols = [c for c, t in profile.items() if t == "categorical"]
    date_cols = [c for c, t in profile.items() if t == "datetime"]
    doc.add_paragraph(
        f"The file had {len(raw_df):,} rows and {raw_df.shape[1]} columns. After automated cleaning, "
        f"{metrics['rows_after']:,} rows and {metrics['cols_kept']} columns remain. "
        f"{metrics['filled']:,} missing values were filled and {metrics['flagged']:,} unusual values were "
        f"flagged for review (flagged values are kept, not deleted). The data has {len(num_cols)} numeric, "
        f"{len(cat_cols)} category and {len(date_cols)} date columns.")
    for n in notes or []:
        doc.add_paragraph(n, style="List Bullet")

    # ---- cleaning ----
    doc.add_heading("2. What the cleaning changed", level=1)
    _table(doc, ["Measure", "Value"], [
        ["Rows after cleaning", f"{metrics['rows_after']:,}"],
        ["Columns kept", f"{metrics['cols_kept']} of {metrics['cols_raw']}"],
        ["Rows removed (duplicates or empty)", f"{metrics['rows_removed']:,}"],
        ["Missing values filled", f"{metrics['filled']:,}"],
        ["Unusual values flagged (1.5 × IQR rule)", f"{metrics['flagged']:,}"],
    ], widths=[3.5, 2.5])
    if cleaning_lines:
        for line in cleaning_lines:
            doc.add_paragraph(line, style="List Bullet")
    else:
        doc.add_paragraph("No cleaning was needed; the data was already tidy.")
    doc.add_heading("Column types detected", level=2)
    _table(doc, ["Column", "Type"], [[c, t] for c, t in profile.items()], widths=[3.5, 2.5])

    # ---- numeric columns ----
    if num_cols:
        doc.add_heading("3. Number columns", level=1)
        doc.add_paragraph("The median is the typical value: half the rows are below it and half above. "
                          "The middle half of rows fall between the 25th and 75th percentiles. When the "
                          "average is far above the median, a few very large values are pulling it up.")
        for col in num_cols[:MAX_NUMERIC]:
            s = clean_df[col].dropna().astype(float)
            if s.empty:
                continue
            doc.add_heading(col, level=2)
            q1, q3 = s.quantile(0.25), s.quantile(0.75)
            _table(doc, ["Rows", "Median", "Middle half", "Average", "Smallest", "Largest"], [[
                f"{len(s):,}", _fmt(float(s.median()), col), f"{_fmt(float(q1), col)} to {_fmt(float(q3), col)}",
                _fmt(float(s.mean()), col), _fmt(float(s.min()), col), _fmt(float(s.max()), col)]])
            img, log = _numeric_chart(s, col)
            doc.add_picture(img, width=Inches(6.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            method, why = d.choose_binning(s)
            if log:
                doc.add_paragraph(f"Shown on a log scale because {why}.").runs[0].italic = True
            ratio = s.mean() / s.median() if s.median() else np.nan
            if np.isfinite(ratio) and ratio > 3:
                doc.add_paragraph(f"The average is about {ratio:,.0f} times the median, so the median is the better "
                                  "description of a typical row.")
        if len(num_cols) > MAX_NUMERIC:
            doc.add_paragraph(f"{len(num_cols) - MAX_NUMERIC} more number columns are not shown here.")

        if len(num_cols) >= 2:
            corr = clean_df[num_cols].corr(method="spearman")
            pairs = []
            for i, a in enumerate(num_cols):
                for b in num_cols[i + 1:]:
                    v = corr.loc[a, b]
                    if pd.notna(v):
                        pairs.append((abs(v), a, b, v))
            pairs.sort(reverse=True)
            if pairs:
                doc.add_heading("How the number columns relate", level=2)
                doc.add_paragraph("Rank correlation from −1 to 1: near 0 means no relationship, near 1 means they "
                                  "rise together, near −1 means one falls as the other rises.")
                strength = lambda v: "strong" if abs(v) >= 0.7 else "moderate" if abs(v) >= 0.4 else "weak"
                _table(doc, ["Column", "Column", "Correlation", "Strength"],
                       [[a, b, f"{v:.2f}", strength(v)] for _, a, b, v in pairs[:8]])

    # ---- categorical columns ----
    if cat_cols:
        doc.add_heading("4. Category columns", level=1)
        for col in cat_cols[:MAX_CATEGORICAL]:
            counts = clean_df[col].astype(str).value_counts()
            if counts.empty:
                continue
            doc.add_heading(col, level=2)
            n = counts.sum()
            doc.add_paragraph(f"{counts.size:,} different values. The most common is "
                              f"“{counts.index[0]}” ({counts.iloc[0] / n:.0%} of rows).")
            _table(doc, ["Value", "Rows", "Share"],
                   [[str(k)[:60], f"{v:,}", f"{v / n:.1%}"] for k, v in counts.head(10).items()],
                   widths=[3.8, 1.1, 1.1])
            doc.add_picture(_categorical_chart(counts, col), width=Inches(6.2))
        if len(cat_cols) > MAX_CATEGORICAL:
            doc.add_paragraph(f"{len(cat_cols) - MAX_CATEGORICAL} more category columns are not shown here.")

    # ---- dates ----
    if date_cols:
        doc.add_heading("5. Date columns", level=1)
        rows = []
        for col in date_cols:
            s = pd.to_datetime(clean_df[col], errors="coerce").dropna()
            if not s.empty:
                rows.append([col, f"{s.min():%Y-%m-%d}", f"{s.max():%Y-%m-%d}", f"{len(s):,}"])
        if rows:
            _table(doc, ["Column", "Earliest", "Latest", "Rows with a date"], rows)

    doc.add_heading("Method notes", level=1)
    for line in [
        "Missing numbers are filled with the column median and missing categories with the most common value; "
        "columns that are mostly empty are set aside instead of filled.",
        "Unusual values are flagged with the 1.5 × interquartile-range rule and kept in the data.",
        "Identifier columns (such as award or record numbers) are removed from analysis because every value is unique.",
        "Code columns such as NAICS, ZIP and PSC are treated as categories, not amounts.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
